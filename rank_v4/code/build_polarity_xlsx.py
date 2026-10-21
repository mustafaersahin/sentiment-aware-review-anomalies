"""Builds the annotator workbook from the id+text CSV (system python with openpyxl).
Usage: python3 build_polarity_xlsx.py <sheet_csv> <out_xlsx>"""
import sys, csv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding="utf-8")))
wb = Workbook(); g = wb.active; g.title = "Yönerge"
lines = [
    ("Metin yönü etiketleme", True),
    ("", False),
    ("Görev: Her satırda bir ürün yorumu metni var. Metnin genel yönünü işaretleyin.", False),
    ("", False),
    ("Etiketler (Etiketleme sayfasındaki C sütunu, açılır listeden seçin):", True),
    ("olumlu – Metin açıkça olumlu bir görüş bildiriyor.", False),
    ("olumsuz – Metin açıkça olumsuz bir görüş bildiriyor.", False),
    ("belirsiz – Metnin yönü belirgin değil (ör. olumlu ve olumsuz yönler dengede, anlam çıkarılamıyor ya da metin Türkçe değil).", False),
    ("", False),
    ("Kurallar:", True),
    ("1. Her metin için tek bir etiket seçin; etiketledikten sonra yeniden etiketleme turu yapmayın.", False),
    ("2. Yalnızca metnin kendisine bakın; konunun ürün, kargo ya da satıcı olması fark etmez, metnin genel yönüne karar verin.", False),
    ("3. Metin '…Devamını Oku' ile kesilmişse görünen kısma göre karar verin.", False),
    ("4. Satırların sırası rastgeledir; başka bir kaynaktan bilgi aramayın.", False),
    ("5. Kararsız kalırsanız 'belirsiz' seçin.", False),
    ("", False),
    ("Düzenlenecek hücreler: yalnızca 'Etiketleme' sayfasındaki sarı C sütunu.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = g.cell(row=i, column=1, value=t); c.font = Font(bold=b, size=13 if i == 1 else 11); c.alignment = Alignment(wrap_text=True)
g.column_dimensions["A"].width = 110
r = len(lines) + 2
g.cell(row=r, column=1, value="İlerleme").font = Font(bold=True)
g.cell(row=r + 1, column=1, value="Etiketlenen metin sayısı:"); g.cell(row=r + 1, column=2, value=f"=COUNTA(Etiketleme!C2:C{len(rows) + 1})")
g.cell(row=r + 2, column=1, value="Toplam metin sayısı:"); g.cell(row=r + 2, column=2, value=f"=COUNTA(Etiketleme!A2:A{len(rows) + 1})")
e = wb.create_sheet("Etiketleme")
for j, h in enumerate(("ID", "Metin", "Etiket"), 1):
    c = e.cell(row=1, column=j, value=h); c.font = Font(bold=True)
yellow = PatternFill("solid", fgColor="FFF2CC")
dv = DataValidation(type="list", formula1='"olumlu,olumsuz,belirsiz"', allow_blank=True, showErrorMessage=True,
                    errorTitle="Geçersiz etiket", error="Lütfen listeden seçin: olumlu, olumsuz, belirsiz")
e.add_data_validation(dv)
for i, row in enumerate(rows, 2):
    e.cell(row=i, column=1, value=row["id"])
    c = e.cell(row=i, column=2, value=row["text"]); c.alignment = Alignment(wrap_text=True, vertical="top")
    t = e.cell(row=i, column=3); t.fill = yellow; dv.add(t)
e.column_dimensions["A"].width = 9; e.column_dimensions["B"].width = 90; e.column_dimensions["C"].width = 12
e.freeze_panes = "A2"
wb.save(out)
